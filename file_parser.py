from docx import Document
import PyPDF2

def read_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    return "\n".join(full_text)

def read_pdf(file_path):
    text = ""
    with open(file_path, "rb") as f:
        pdf_reader = PyPDF2.PdfReader(f)
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text.strip()

# 文件分发读取
def parse_upload_file(file):
    text = ""
    filename = file.name
    if filename.endswith(".pdf"):
        pdf_reader = PyPDF2.PdfReader(file)
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text
    elif filename.endswith(".docx"):
        doc = Document(file)
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text.strip()